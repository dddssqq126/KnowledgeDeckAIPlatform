import base64
import io

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from app.db.models import FileStatus, IngestionMode
from app.features.rag.services import document_parser, image_ingestion, ingestion


@pytest.fixture(scope="session", autouse=True)
def _run_migrations() -> None:
    pass


@pytest.fixture(autouse=True)
def _patch_app_db() -> None:
    pass


@pytest.fixture(autouse=True)
def _patch_app_storage() -> None:
    pass


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _deck_with_picture_and_shape() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    text_box.text = "季度營收趨勢"
    slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(1),
        Inches(2),
        Inches(1),
        Inches(1),
    )
    slide.shapes.add_picture(io.BytesIO(PNG_1X1), Inches(2), Inches(2))
    slide.shapes.add_picture(io.BytesIO(PNG_1X1), Inches(3), Inches(2))
    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()


def test_extract_pptx_images_only_keeps_pictures_and_deduplicates() -> None:
    images = document_parser.extract_pptx_images(_deck_with_picture_and_shape())

    assert len(images) == 1
    assert images[0].extension == "png"
    assert images[0].content_type == "image/png"
    assert images[0].page_number == 1
    assert images[0].image_index == 1
    assert images[0].page_text == "季度營收趨勢"
    assert images[0].data == PNG_1X1


def test_extract_pdf_images_preserves_page_and_deduplicates(monkeypatch) -> None:
    class FakeImage:
        name = "figure.png"
        data = PNG_1X1

    class FakePage:
        images = [FakeImage()]

        def __init__(self, text: str) -> None:
            self._text = text

        def extract_text(self) -> str:
            return self._text

    class FakeReader:
        def __init__(self, _stream) -> None:
            self.pages = [FakePage("page one"), FakePage("page two")]

    monkeypatch.setattr(document_parser, "PdfReader", FakeReader)

    images = document_parser.extract_pdf_images(b"fake pdf")

    assert len(images) == 1
    assert images[0].data == PNG_1X1
    assert images[0].extension == "png"
    assert images[0].content_type == "image/png"
    assert images[0].page_number == 1
    assert images[0].image_index == 1
    assert images[0].page_text == "page one"


@pytest.mark.asyncio
async def test_image_mode_skips_document_parser(monkeypatch) -> None:
    class FakeSession:
        commits = 0

        async def commit(self) -> None:
            self.commits += 1

    row = type(
        "FileRow",
        (),
        {
            "ingestion_mode": IngestionMode.IMAGE,
            "extension": "pdf",
            "status": FileStatus.UPLOADED,
            "status_error": None,
        },
    )()

    async def fake_images(**_kwargs) -> int:
        return 2

    def fail_parse(*_args, **_kwargs):
        raise AssertionError("document parser must not run in image mode")

    monkeypatch.setattr(image_ingestion, "ingest_document_images", fake_images)
    monkeypatch.setattr(document_parser, "parse", fail_parse)
    session = FakeSession()

    await ingestion.ingest_file(session=session, file_row=row, data=b"pdf")

    assert row.status is FileStatus.INDEXED
    assert row.status_error is None
    assert session.commits == 1


@pytest.mark.asyncio
async def test_generate_label_falls_back_when_llm_returns_invalid_json(
    monkeypatch,
) -> None:
    class BrokenClient:
        def __init__(self, **_kwargs) -> None:
            pass

        async def create_chat_completion(self, *_args, **_kwargs):
            return {"choices": [{"message": {"content": "not json"}}]}

    monkeypatch.setattr(image_ingestion, "ChatModelClient", BrokenClient)
    file_row = type(
        "FileRow",
        (),
        {"id": 7, "filename": "sales.pptx"},
    )()
    image = document_parser.extract_pptx_images(_deck_with_picture_and_shape())[0]

    label = await image_ingestion.generate_label(file_row, image)

    assert label.name == "sales.pptx－第 1 頁－圖片 1"
    assert label.description == "季度營收趨勢"
