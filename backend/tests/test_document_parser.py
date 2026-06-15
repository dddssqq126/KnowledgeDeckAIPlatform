import io
from datetime import date

import pytest
from openpyxl import Workbook

from app.features.knowledge_bases.services import file_service
from app.features.rag.services import document_parser


def _xlsx_bytes() -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Vendor", "Teradyne"])
    summary.append(["Platform", "UltraFLEX"])
    summary.append(["Enabled", True])
    summary.append(["Launched", date(2026, 6, 11)])
    summary.append(["Formula", "=SUM(1,2)"])

    wb.create_sheet("Empty")

    details = wb.create_sheet("Details")
    details.append(["Limit", 42])
    details.append(["Inline note", "OK"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_xlsx_extracts_non_empty_worksheets() -> None:
    segments = document_parser.parse("xlsx", _xlsx_bytes())

    assert [segment.page_number for segment in segments] == [1, 3]
    assert segments[0].text == (
        "Sheet: Summary\n"
        "Vendor\tTeradyne\n"
        "Platform\tUltraFLEX\n"
        "Enabled\tTRUE\n"
        "Launched\t2026-06-11T00:00:00\n"
        "Formula\t=SUM(1,2)"
    )
    assert segments[1].text == "Sheet: Details\nLimit\t42\nInline note\tOK"


def test_parse_csv_utf8_handles_quotes_and_blank_rows() -> None:
    data = b'Name,Value\n"ACME, Inc",42\n\nFeature,Enabled\n'

    segments = document_parser.parse("csv", data)

    assert [segment.text for segment in segments] == [
        "Name\tValue\nACME, Inc\t42\nFeature\tEnabled"
    ]


def test_parse_csv_utf8_bom() -> None:
    data = "Name,Value\n平台,UltraFLEX\n".encode("utf-8-sig")

    segments = document_parser.parse("csv", data)

    assert segments[0].text == "Name\tValue\n平台\tUltraFLEX"


def test_parse_csv_cp950() -> None:
    data = "名稱,值\n平台,測試\n".encode("cp950")

    segments = document_parser.parse("csv", data)

    assert segments[0].text == "名稱\t值\n平台\t測試"


def test_parse_csv_sniffs_tab_delimiter() -> None:
    data = b"Name\tValue\nPlatform\tUltraFLEX\n"

    segments = document_parser.parse("csv", data)

    assert segments[0].text == "Name\tValue\nPlatform\tUltraFLEX"


def test_file_service_accepts_xlsx_ooxml_upload() -> None:
    data = _xlsx_bytes()

    assert file_service.validate_extension("workbook.xlsx") == "xlsx"
    file_service.validate_content("xlsx", data[:1024])


def test_file_service_accepts_csv_upload() -> None:
    data = "名稱,值\n平台,測試\n".encode("cp950")

    assert file_service.validate_extension("table.csv") == "csv"
    file_service.validate_content("csv", data[:1024])


def test_file_service_rejects_malformed_xlsx_upload() -> None:
    with pytest.raises(file_service.ValidationError) as exc:
        file_service.validate_content("xlsx", b"not a zip")

    assert exc.value.code == "invalid_content"


def test_file_service_rejects_binary_csv_upload() -> None:
    with pytest.raises(file_service.ValidationError) as exc:
        file_service.validate_content("csv", b"name,value\x00still-binary")

    assert exc.value.code == "invalid_content"
